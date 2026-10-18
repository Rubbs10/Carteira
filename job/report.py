"""Relatório mensal (PDF + Excel) e alertas por email.

Uso:  python job/report.py monthly   (relatório completo com anexos)
      python job/report.py alerts    (só envia se houver avisos novos)
Variáveis: SUPABASE_URL, SUPABASE_SERVICE_KEY, REPORT_EMAIL (o teu email), RESEND_API_KEY.
DRY_RUN=pasta  -> não envia, grava os ficheiros nessa pasta (para testar).
"""
import base64
import io
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone

import update_prices as up

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
DRIFT_PP = 5.0     # alerta quando uma divisão se afasta da meta mais do que isto (pontos percentuais)
DROP_PCT = 7.0     # alerta quando um ativo cai mais do que isto num dia
STALE_DAYS = 4     # alerta quando um preço não é atualizado há mais dias do que isto


def eur(v, sign=False):
    s = f"{abs(v):,.2f}".replace(",", " ").replace(".", ",")
    pre = ("+" if v > 0 else "-" if v < 0 else "") if sign else ("-" if v < 0 else "")
    return f"{pre}{s} €"


def pct(v, sign=True):
    s = f"{abs(v) * 100:.2f}".replace(".", ",")
    return (("+" if v > 0 else "-" if v < 0 else "") if sign else "") + s + " %"


# ───────── dados ─────────
def find_user(email):
    """Procura a conta pelo email; se não existir, usa o único utilizador com posições (app pessoal)."""
    try:
        req = urllib.request.Request(up.URL + "/auth/v1/admin/users?per_page=200",
                                     headers={"apikey": up.KEY, "Authorization": "Bearer " + up.KEY})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = json.loads(r.read())
        for u in data.get("users", data if isinstance(data, list) else []):
            if (u.get("email") or "").strip().lower() == email.strip().lower():
                return u["id"]
    except Exception as e:
        print(f"  lista de utilizadores indisponível: {e}")
    ids = sorted({p["user_id"] for p in (up.api("GET", "positions?select=user_id") or [])})
    if len(ids) == 1:
        print("  email do segredo não coincide com a conta; usei o único utilizador com posições")
        return ids[0]
    raise SystemExit(f"Utilizador não encontrado ({len(ids)} contas com posições). Confirma o segredo REPORT_EMAIL.")


def load(uid):
    q = lambda path: up.api("GET", path) or []
    d = {"divisions": q(f"divisions?user_id=eq.{uid}&select=*&order=sort_order"),
         "positions": q(f"positions?user_id=eq.{uid}&select=*"),
         "snaps": q(f"snapshots?user_id=eq.{uid}&select=*&order=month")}
    syms = sorted({p["symbol"] for p in d["positions"] if p.get("symbol")})
    quotes = {}
    for i in range(0, len(syms), 80):
        part = ",".join('"' + s.replace('"', "") + '"' for s in syms[i:i + 80])
        for row in q("quotes?select=*&symbol=in.(" + urllib.parse.quote(part) + ")"):
            quotes[row["symbol"]] = row
    d["quotes"] = quotes
    fx = q("fx_rates?select=*&pair=eq.USD_EUR")
    d["fx"] = float(fx[0]["rate"]) if fx else None
    try:
        since = (date.today() - timedelta(days=45)).isoformat()
        d["txs"] = q(f"transactions?user_id=eq.{uid}&select=kind,occurred_at,ticker,amount,currency&occurred_at=gte.{since}&order=occurred_at")
    except RuntimeError:
        d["txs"] = []
    return d


# ───────── cálculo (função pura, testável) ─────────
def position_price(p, quotes, today):
    if p.get("manual"):
        price = p.get("manual_price")
        if price is None:
            return None, None, None
        price = float(price)
        if p.get("accrual_rate") is not None and p.get("accrual_start"):
            days = max(0, (today - date.fromisoformat(str(p["accrual_start"])[:10])).days)
            price *= (1 + float(p["accrual_rate"]) / 100) ** (days / 365)
        return price, price, None
    q = quotes.get(p.get("symbol"))
    if not q or (q.get("currency") and str(q["currency"]).upper() != p["currency"]):
        return None, None, None
    prev = float(q["prev_close"]) if q.get("prev_close") else None
    return float(q["price"]), prev, q.get("as_of")


def build(d, today=None, now=None):
    today = today or date.today()
    now = now or datetime.now(timezone.utc)
    fx = d["fx"]
    divs = {x["id"]: {"id": x["id"], "name": x["name"], "target": (float(x["target_pct"]) if x.get("target_pct") is not None else None),
                      "value": 0.0, "invested": 0.0} for x in d["divisions"]}
    divs["_none"] = {"id": "_none", "name": "Sem divisão", "target": None, "value": 0.0, "invested": 0.0}
    rows, missing, alerts = [], [], []
    for p in d["positions"]:
        price, prev, as_of = position_price(p, d["quotes"], today)
        r = fx if p["currency"] == "USD" else 1.0
        if price is None or r is None:
            missing.append(p["ticker"])
            continue
        qty, avg = float(p["qty"]), float(p["avg_cost"])
        value, cost = qty * price * r, qty * avg * r
        day = (price / prev - 1) if prev else 0.0
        key = p.get("division_id") if p.get("division_id") in divs else "_none"
        divs[key]["value"] += value
        divs[key]["invested"] += cost
        rows.append({"ticker": p["ticker"], "name": p.get("name") or "", "division": divs[key]["name"], "qty": qty, "avg": avg,
                     "price": price, "ccy": p["currency"], "value": value, "cost": cost, "pl": value - cost,
                     "plpct": (value / cost - 1) if cost else 0.0, "day": day, "manual": bool(p.get("manual"))})
        if day <= -DROP_PCT / 100:
            alerts.append((f"drop:{p['ticker']}", f"{p['ticker']} caiu {pct(day, False)} no último dia."))
        if as_of and not p.get("manual"):
            age = (now - datetime.fromisoformat(str(as_of).replace("Z", "+00:00"))).days
            if age >= STALE_DAYS and now.weekday() < 5:
                alerts.append((f"stale:{p['ticker']}", f"O preço de {p['ticker']} não é atualizado há {age} dias."))
    rows.sort(key=lambda x: -x["value"])
    total = sum(x["value"] for x in rows)
    invested = sum(x["cost"] for x in rows)
    dl = [v for v in divs.values() if v["value"] > 0 or v["target"]]
    for v in dl:
        v["weight"] = v["value"] / total if total else 0.0
        v["drift"] = (v["weight"] * 100 - v["target"]) if v["target"] is not None else None
        if v["drift"] is not None and abs(v["drift"]) >= DRIFT_PP:
            sentido = "acima" if v["drift"] > 0 else "abaixo"
            alerts.append((f"drift:{v['id']}", f"{v['name']} está {abs(v['drift']):.1f} pontos {sentido} da meta ({v['weight'] * 100:.1f} % em vez de {v['target']:.0f} %)."))
    if missing:
        alerts.append(("missing:" + ",".join(sorted(missing)), "Sem preço para: " + ", ".join(sorted(missing)) + "."))
    # variação do mês face ao fim do mês anterior
    ref_day = today - timedelta(days=1) if today.day == 1 else today
    month_first = ref_day.replace(day=1)
    prev_first = (month_first - timedelta(days=1)).replace(day=1)
    ref = next((s for s in d["snaps"] if str(s["month"])[:7] == prev_first.isoformat()[:7]), None)
    chg = None
    if ref:
        dv, di = total - float(ref["value"]), invested - float(ref["invested"])
        chg = {"value": dv, "contrib": di, "market": dv - di, "since": prev_first}
    # dividendos e juros dos últimos 31 dias
    cut = now - timedelta(days=31)
    inc = 0.0
    for t in d["txs"]:
        if t["kind"] in ("dividend", "interest") and datetime.fromisoformat(str(t["occurred_at"]).replace("Z", "+00:00")) >= cut:
            a = float(t["amount"]) * (fx if (t["currency"] == "USD" and fx) else 1.0)
            inc += a
    return {"total": total, "invested": invested, "pl": total - invested, "rows": rows, "divs": sorted(dl, key=lambda x: -x["value"]),
            "alerts": alerts, "chg": chg, "income": inc, "missing": missing, "month": ref_day, "snaps": d["snaps"], "today": today}


# ───────── ficheiros ─────────
def make_xlsx(R):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0A6A8A")
    E = '#,##0.00 "€";[Red]-#,##0.00 "€"'

    def sheet(ws, headers, data, fmts, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
        for i, f in enumerate(fmts, 1):
            for c in ws[get_column_letter(i)][1:]:
                if f:
                    c.number_format = f
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 30

    ws = wb.active
    ws.title = "Resumo"
    m = R["month"]
    ws.append([f"PatrimÓ - {MESES[m.month - 1]} de {m.year}"])
    ws["A1"].font = Font(bold=True, size=14)
    lines = [("Valor total", R["total"]), ("Total investido", R["invested"]), ("Ganho total", R["pl"])]
    if R["chg"]:
        lines += [("Variação desde o mês anterior", R["chg"]["value"]), ("  das quais entradas de dinheiro", R["chg"]["contrib"]), ("  das quais mercado", R["chg"]["market"])]
    lines.append(("Dividendos e juros (31 dias)", R["income"]))
    ws.append([])
    for k, v in lines:
        ws.append([k, v])
        ws.cell(ws.max_row, 2).number_format = E
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    sheet(wb.create_sheet("Divisões"), ["Divisão", "Valor (€)", "Peso %", "Meta %", "Desvio (p.p.)"],
          [[v["name"], round(v["value"], 2), round(v["weight"] * 100, 2), v["target"], round(v["drift"], 2) if v["drift"] is not None else None] for v in R["divs"]],
          [None, E, "0.00", "0.00", "+0.00;-0.00"], [26, 16, 10, 10, 14])
    sheet(wb.create_sheet("Posições"), ["Ativo", "Nome", "Divisão", "Qtd.", "Preço", "Moeda", "Valor (€)", "Custo (€)", "Ganho (€)", "Ganho %", "Dia %"],
          [[x["ticker"], x["name"], x["division"], x["qty"], round(x["price"], 4), x["ccy"], round(x["value"], 2), round(x["cost"], 2), round(x["pl"], 2), x["plpct"], x["day"]] for x in R["rows"]],
          [None, None, None, "#,##0.######", "#,##0.00", None, E, E, E, "0.00%", "0.00%"], [12, 30, 18, 10, 11, 8, 14, 14, 14, 10, 10])
    sheet(wb.create_sheet("Evolução"), ["Mês", "Valor (€)", "Investido (€)", "Ganho (€)"],
          [[str(s["month"])[:7], float(s["value"]), float(s["invested"]), float(s["value"]) - float(s["invested"])] for s in R["snaps"]],
          [None, E, E, E], [12, 16, 16, 16])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pdf(R):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    acc, ink, mut, line, gain, loss = colors.HexColor("#0a6a8a"), colors.HexColor("#0e1a24"), colors.HexColor("#586a78"), colors.HexColor("#d9e1e7"), colors.HexColor("#0a7a4e"), colors.HexColor("#bf3a2c")
    ss = getSampleStyleSheet()
    H1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=ink, alignment=0, spaceAfter=2)
    H2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=acc, spaceBefore=14, spaceAfter=6)
    P = ParagraphStyle("p", parent=ss["BodyText"], fontName="Helvetica", fontSize=9, textColor=ink, leading=12)
    S = ParagraphStyle("s", parent=P, fontSize=8, textColor=mut)
    m = R["month"]
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="PatrimÓ", author="PatrimÓ")
    el = [Paragraph("PatrimÓ", H1), Paragraph(f"Relatório de {MESES[m.month - 1]} de {m.year}", S), Spacer(1, 8)]

    def kpi(label, val, color=ink):
        return [Paragraph(f'<font size="15" color="{color.hexval().replace("0x", "#")}"><b>{val}</b></font>', P), Paragraph(label, S)]
    cells = [kpi("Valor total", eur(R["total"])), kpi("Total investido", eur(R["invested"])), kpi("Ganho total", eur(R["pl"], True), gain if R["pl"] >= 0 else loss)]
    if R["chg"]:
        cells.append(kpi("Variação do mês", eur(R["chg"]["value"], True), gain if R["chg"]["value"] >= 0 else loss))
    t = Table([[c[0] for c in cells], [c[1] for c in cells]], colWidths=[(178 * mm) / len(cells)] * len(cells))
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), 0.6, line), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    el += [t]
    if R["chg"]:
        c = R["chg"]
        el += [Spacer(1, 4), Paragraph(f"Desde o fim do mês anterior: {eur(c['contrib'], True)} de entradas de dinheiro e {eur(c['market'], True)} de ganho do mercado. Dividendos e juros recebidos nos últimos 31 dias: {eur(R['income'])}.", S)]
    elif R["income"]:
        el += [Spacer(1, 4), Paragraph(f"Dividendos e juros recebidos nos últimos 31 dias: {eur(R['income'])}.", S)]

    def table(head, data, widths, right_from=1, extra=None):
        T = Table([head] + data, colWidths=[w * mm for w in widths], repeatRows=1)
        st = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), acc),
              ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5), ("TEXTCOLOR", (0, 1), (-1, -1), ink), ("ALIGN", (right_from, 0), (-1, -1), "RIGHT"),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f4f6")]), ("LINEBELOW", (0, -1), (-1, -1), 0.5, line),
              ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        T.setStyle(TableStyle(st + (extra or [])))
        return T

    el += [Paragraph("Divisões e metas", H2)]
    dd = [[v["name"], eur(v["value"]), pct(v["weight"], False), (f"{v['target']:.0f} %" if v["target"] is not None else "-"), (f"{v['drift']:+.1f} p.p.".replace(".", ",") if v["drift"] is not None else "-")] for v in R["divs"]]
    ex = []
    for i, v in enumerate(R["divs"], 1):
        if v["drift"] is not None and abs(v["drift"]) >= DRIFT_PP:
            ex.append(("TEXTCOLOR", (4, i), (4, i), loss))
    el += [table(["Divisão", "Valor", "Peso", "Meta", "Desvio"], dd, [62, 36, 26, 26, 28], extra=ex)]
    el += [Paragraph("Maiores posições", H2)]
    pr = [[x["ticker"], eur(x["value"]), eur(x["pl"], True), pct(x["plpct"]), pct(x["day"])] for x in R["rows"][:12]]
    ex = []
    for i, x in enumerate(R["rows"][:12], 1):
        ex.append(("TEXTCOLOR", (2, i), (3, i), gain if x["pl"] >= 0 else loss))
        ex.append(("TEXTCOLOR", (4, i), (4, i), gain if x["day"] >= 0 else loss))
    el += [table(["Ativo", "Valor", "Ganho", "Ganho %", "Dia"], pr, [50, 36, 36, 28, 28], extra=ex)]
    el += [Paragraph("Avisos", H2)]
    if R["alerts"]:
        el += [Paragraph("• " + a[1], P) for a in R["alerts"]]
    else:
        el += [Paragraph("Sem avisos. Tudo dentro dos limites que definiste.", P)]
    el += [Spacer(1, 14), Paragraph("Valores em euros, calculados com as últimas cotações e o câmbio atual. Poupança com juro: estimativa bruta, sem impostos. Relatório informativo, não é aconselhamento financeiro.", S)]
    doc.build(el)
    return out.getvalue()


def make_html(R, title):
    lis = "".join(f"<li>{a[1]}</li>" for a in R["alerts"]) or "<li>Sem avisos.</li>"
    chg = f"<p>Variação do mês: <b>{eur(R['chg']['value'], True)}</b></p>" if R["chg"] else ""
    return (f"<div style='font-family:Segoe UI,Arial,sans-serif;color:#0e1a24;max-width:560px'><h2 style='color:#0a6a8a;margin:0 0 8px'>{title}</h2>"
            f"<p style='font-size:22px;margin:4px 0'><b>{eur(R['total'])}</b></p><p>Ganho total: <b>{eur(R['pl'], True)}</b></p>{chg}"
            f"<h3 style='margin:16px 0 4px'>Avisos</h3><ul>{lis}</ul><p style='color:#586a78;font-size:12px'>PDF e Excel completos em anexo.</p></div>")


def send(to, subject, html, files):
    out = os.environ.get("DRY_RUN")
    if out:
        os.makedirs(out, exist_ok=True)
        for name, content in files:
            open(os.path.join(out, name), "wb").write(content)
        open(os.path.join(out, "email.html"), "w").write(f"<h1>{subject}</h1>{html}")
        print("DRY_RUN:", subject, [n for n, _ in files])
        return
    body = {"from": "PatrimÓ <onboarding@resend.dev>", "to": [to], "subject": subject, "html": html,
            "attachments": [{"filename": n, "content": base64.b64encode(c).decode()} for n, c in files]}
    req = urllib.request.Request("https://api.resend.com/emails", data=json.dumps(body).encode(), method="POST",
                                 headers={"Authorization": "Bearer " + os.environ["RESEND_API_KEY"], "Content-Type": "application/json", "User-Agent": "patrimo-report/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            print("email enviado:", r.status)
    except urllib.error.HTTPError as e:
        raise SystemExit(f"Falha ao enviar email: {e.code} {e.read().decode()[:300]}")


def main(mode):
    if not up.URL or not up.KEY:
        sys.exit("Faltam SUPABASE_URL e/ou SUPABASE_SERVICE_KEY")
    to = os.environ.get("REPORT_EMAIL", "").strip()
    if not to:
        sys.exit("Falta REPORT_EMAIL")
    if not os.environ.get("DRY_RUN") and not os.environ.get("RESEND_API_KEY"):
        sys.exit("Falta RESEND_API_KEY")
    uid = find_user(to)
    R = build(load(uid))
    m = R["month"]
    if mode == "monthly":
        title = f"PatrimÓ - relatório de {MESES[m.month - 1]} de {m.year}"
        send(to, title, make_html(R, title), [(f"patrimo-{m.year}-{m.month:02d}.pdf", make_pdf(R)), (f"patrimo-{m.year}-{m.month:02d}.xlsx", make_xlsx(R))])
        return
    try:
        sent = {r["key"]: r["sent_on"] for r in up.api("GET", f"alerts_sent?user_id=eq.{uid}&select=key,sent_on") or []}
    except RuntimeError as e:
        print(f"alertas ignorados (corre o phase4.sql no Supabase): {e}")
        return
    cutoff = (date.today() - timedelta(days=7)).isoformat()
    new = [a for a in R["alerts"] if a[0] not in sent or sent[a[0]] < cutoff]
    keys_now = {a[0] for a in R["alerts"]}
    for k in list(sent):  # quando o problema desaparece, o aviso pode voltar a ser enviado no futuro
        if k not in keys_now:
            up.api("DELETE", f"alerts_sent?user_id=eq.{uid}&key=eq." + urllib.parse.quote(k, safe=""), prefer="return=minimal")
    if not new:
        print("sem avisos novos")
        return
    R2 = dict(R, alerts=new)
    title = "PatrimÓ - " + ("1 aviso" if len(new) == 1 else f"{len(new)} avisos")
    send(to, title, make_html(R2, title).replace("PDF e Excel completos em anexo.", ""), [])
    up.api("POST", "alerts_sent", [{"user_id": uid, "key": a[0], "sent_on": date.today().isoformat()} for a in new],
           prefer="resolution=merge-duplicates,return=minimal")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "alerts")
